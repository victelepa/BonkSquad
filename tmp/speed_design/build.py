from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[2]
doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Inches(8.5), Inches(11)
s.top_margin = s.bottom_margin = Inches(.7)
s.left_margin = s.right_margin = Inches(.8)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    st = doc.styles[name]
    st.font.name = 'Calibri'
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Microsoft YaHei')
normal = doc.styles['Normal']
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.12
for name, size in [('Title',24), ('Subtitle',11), ('Heading 1',14), ('Heading 2',11)]:
    doc.styles[name].font.size = Pt(size)
    doc.styles[name].paragraph_format.space_before = Pt(10 if name.startswith('Heading') else 0)
    doc.styles[name].paragraph_format.space_after = Pt(6)

def p(text='', style=None):
    return doc.add_paragraph(text, style)
def h(text):
    doc.add_heading(text, 1)
def blank():
    for _ in range(1):
        x = p()
        x.paragraph_format.space_after = Pt(9)
def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for c, w, txt in zip(t.rows[0].cells, widths, headers):
        c.width = Inches(w)
        c.text = txt
    for row in rows:
        for c,w,txt in zip(t.add_row().cells,widths,row):
            c.width = Inches(w)
            c.text = txt
    for i,row in enumerate(t.rows):
        for c in row.cells:
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr = c._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                el = OxmlElement('w:'+edge)
                for key,val in [('val','single'),('sz','4'),('color','D9D9D9')]:
                    el.set(qn('w:'+key),val)
                borders.append(el)
            pr.append(borders)
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), 'E5EDF3' if i == 0 else ('F7F9FB' if i%2==0 else 'FFFFFF'))
            pr.append(shade)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge); el.set(qn('w:w'),'90'); el.set(qn('w:type'),'dxa'); margins.append(el)
            pr.append(margins)
            for para in c.paragraphs:
                para.paragraph_format.space_after = Pt(1)
                para.paragraph_format.line_spacing = 1.05
                for r in para.runs:
                    r.font.size = Pt(10)
                    r.bold = i == 0
        cant = OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(cant)
    rep = OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(rep)
    p().paragraph_format.space_after = Pt(0)

p('Speed System Design', 'Title')
p('English Version', 'Subtitle')
p('The speed system shows the player’s current movement speed and whether a collision can instantly kill a small or large enemy. Enemy collisions either instantly kill the enemy or make the player bounce off without damaging the enemy. Kills provide no speed boost.')
h('1 Speed Display')
p('Display a horizontal speed bar in the lower-left corner of the screen, similar to a racing game speedometer. The fill amount updates with the player’s current speed: empty at zero and full at the configured maximum speed. Use the same scale in ball and robot forms so the threshold markers remain consistent.')
p('Use a fixed color treatment, without a speed-dependent gradient or color changes for dizziness. Fill position shows speed and instant-kill eligibility. Keep the bar, labels, and both threshold lines stable and readable; the lines must not shake with surrounding effects. Dizziness effects and warnings are defined separately in Dizziness System Design.')
h('2 Instant Kill Markers')
p('Place two distinct lines or markers within the bar. Label the lower threshold for small enemies and the higher threshold for large enemies. Their positions must correspond to the configured speed values rather than arbitrary visual percentages. The large-enemy threshold must be higher than the small-enemy threshold; both must lie within the displayed range.')
h('3 Enemy Collision Rules')
p('Use the player’s speed at the moment of collision. Reaching a threshold is sufficient, including when speed is exactly equal to that threshold. The speed bar and collision logic must use the same speed value and threshold settings.')
table(['Player speed at collision','Small enemy','Large enemy'],[
('Below the small-enemy threshold','Bounce; no damage','Bounce; no damage'),
('At or above the small-enemy threshold, below the large-enemy threshold','Instant kill','Bounce; no damage'),
('At or above the large-enemy threshold','Instant kill','Instant kill')],[3.7,1.6,1.6])
p('Below the relevant threshold, the player bounces off the enemy without damaging it. The bounce changes movement direction but preserves the player’s speed magnitude. Robot-form deceleration continues independently after the bounce. Enemy health and partial damage are not implemented. Killing either enemy type provides no speed boost.')

doc.add_page_break()
h('4 Robot Form Deceleration')
p('In robot form, speed above robot base speed decreases at a configurable constant rate, measured in speed units per second. Clamp the result at robot base speed. This rate is independent of entry speed; higher entry speeds take longer to slow down. Robot base speed is the form’s normal maximum movement speed and should feel very slow.')
p('If speed is already at or below robot base speed, this deceleration must not increase it. Switching back to ball form immediately stops the deceleration and retains current speed. Re-entering robot form resumes the rule from the current speed.')
p('Instant-kill checks remain active during robot-form deceleration whenever collision speed meets the enemy’s threshold. Update the speed bar continuously. Players can actively steer only in robot form; steering does not cancel this deceleration.')
h('5 Configurable Parameters')
p('Expose these settings to designers for tuning. Keep the two enemy thresholds clearly separated.')
table(['Parameter','Requirement'],[
('Maximum speed','Defines the full scale of the speed bar.'),
('Robot base speed','Normal maximum movement speed in robot form; very slow.'),
('Small-enemy instant-kill threshold','Lower of the two instant-kill thresholds.'),
('Large-enemy instant-kill threshold','Higher than the small-enemy threshold.'),
('Robot-form deceleration rate','Positive constant speed reduction per second.'),
('Acceleration angle tolerance','Maximum allowed angular difference, in degrees, between input and current movement direction.'),
('Speed bar appearance','Fixed color treatment; stable threshold lines and readable fill.')],[3.1,3.8])
h('6 Acceleration Sources')
p('In ball form, use W, A, S, D, W+A, W+D, S+A, or S+D to accelerate along the current movement direction. Acceleration is allowed when the angle between input direction and movement direction is at or below the configured tolerance; outside it, this input provides no acceleration.')
p('Ball-form WASD input does not steer the player. Collisions can still change movement direction. Enemy kills provide no acceleration.')
h('7 Deceleration Sources')
p('Robot form naturally reduces speed at the constant rate in Section 4. Switching back to ball form stops it immediately. Bouncing off an enemy causes no speed loss.')

doc.add_page_break()
p('速度系统设计', 'Title')
p('中文版本', 'Subtitle')
p('速度系统用于显示玩家当前移速，并帮助玩家判断碰撞是否可以秒杀小型或大型敌人。敌人碰撞有两种结果：秒杀敌人，或玩家反弹且敌人不受伤。击杀敌人不提供加速。')
h('1 速度显示')
p('在屏幕左下角显示横向速度条，形式参考赛车游戏的速度表。速度条填充量随玩家当前速度实时更新：速度为零时为空，达到配置的最高速度时填满。球形态与机器人形态使用同一速度刻度，保证秒杀阈值标识保持一致。')
p('速度条使用固定配色，不随速度或眩晕变色，以填充位置显示速度与秒杀条件。主体、标签及两条秒杀刻度线保持稳定，不随外围特效晃动。眩晕特效见眩晕系统设计。')
h('2 秒杀阈值标识')
p('在速度条内部设置两条明显的线或标识，较低的一条标注小型敌人秒杀阈值，较高的一条标注大型敌人秒杀阈值。标识位置应对应实际配置的速度数值，不采用任意固定百分比。大型敌人的秒杀阈值必须高于小型敌人，且两者都应位于速度条显示范围内。')
h('3 敌人碰撞规则')
p('以碰撞瞬间玩家的速度进行判定。达到阈值即可秒杀，恰好等于阈值时也视为满足条件。速度条与碰撞逻辑必须使用相同的速度数值和阈值配置。')
table(['碰撞瞬间的玩家速度','小型敌人','大型敌人'],[
('低于小型敌人秒杀阈值','反弹且无伤害','反弹且无伤害'),
('达到小型敌人秒杀阈值，但低于大型敌人秒杀阈值','秒杀','反弹且无伤害'),
('达到大型敌人秒杀阈值','秒杀','秒杀')],[3.7,1.6,1.6])
p('未达到对应阈值时，玩家从敌人身上反弹，敌人不受伤。反弹改变运动方向，但保持玩家反弹前后的速度大小不变；机器人形态的自然减速仍独立继续生效。当前版本不实现敌人血量和部分伤害。击杀小型或大型敌人均不提供加速。')

doc.add_page_break()
h('4 机器人形态减速')
p('机器人形态下，若当前速度高于机器人基础速度，则按可配置的固定减速度自然减速，单位为每秒减少的速度量，最低降至机器人基础速度。减速度不随切换时的初始速度变化，因此初始速度越高，降速所需时间越长。机器人基础速度即该形态的常态最高移速，表现应很慢。')
p('若当前速度已等于或低于机器人基础速度，这段减速不得反向提高速度。切换回球形态时立即停止减速，保留当前速度；再次进入机器人形态时，从当时的速度重新执行该规则。')
p('机器人减速期间仍可在碰撞速度达到阈值时秒杀敌人，速度条持续更新。玩家只有在机器人形态下可以主动转向；转向不取消自然减速。')
h('5 可配置参数')
p('程序需向策划开放以下设置，供后续调试。两档敌人秒杀阈值应有明确区分。')
table(['参数','要求'],[
('最高速度','决定速度条的满格刻度。'),
('机器人基础速度','机器人形态下的常态最高移速，表现应很慢。'),
('小型敌人秒杀阈值','两档秒杀阈值中较低的一档。'),
('大型敌人秒杀阈值','必须高于小型敌人秒杀阈值。'),
('机器人形态减速度','大于零的固定数值，表示每秒减少的速度量。'),
('加速角度容差','输入方向与当前移动方向之间允许的最大夹角，单位为度。'),
('速度条外观','固定配色，阈值刻度线稳定，填充清晰可读。')],[3.1,3.8])
h('6 加速来源')
p('球形态下，通过 W、A、S、D、W+A、W+D、S+A、S+D 共八个方向的输入沿当前移动方向加速。输入方向与当前移动方向的夹角小于或等于配置的角度容差时允许加速；超出容差时，该输入不提供加速。')
p('球形态的 WASD 输入不使玩家主动转向，碰撞仍可改变运动方向。当前版本击杀敌人不提供加速。')
h('7 减速来源')
p('机器人形态按第 4 节规定的固定减速度自然减速，切换回球形态时立即停止。从敌人身上反弹不损失速度。')

doc.core_properties.title = 'Speed System Design 速度系统设计'
doc.core_properties.subject = 'Player speed display and enemy collision rules'
doc.core_properties.author = ''
for st in doc.styles:
    for el in list(st.element.iter(qn('w:pBdr'))):
        el.getparent().remove(el)
for para in doc.paragraphs:
    for el in list(para._p.iter(qn('w:pBdr'))):
        el.getparent().remove(el)
out = ROOT / 'output' / 'docs' / 'Speed_System_Design_EN_ZH.docx'
doc.save(out)
print(out)
